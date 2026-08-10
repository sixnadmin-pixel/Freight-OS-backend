import openpyxl
import psycopg
from psycopg.rows import dict_row
from dotenv import load_dotenv
import os

load_dotenv()

EXCEL_PATH = r"C:\Users\tharu\Downloads\Customer List.xlsx"
CONN_STR = os.getenv("POSTGRES_CONN_STR")

EMP_ID = 8
ASSIGNED_TO = 7
KYC_COMPLETED = True
BATCH_SIZE = 500


def load_customer_names(path: str) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    names = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        val = row[0]
        if val is not None:
            name = str(val).strip()
            if name:
                names.append(name)
    wb.close()
    return names


def build_batch_insert(names: list[str]) -> tuple[str, dict]:
    """Build a single INSERT with multiple VALUES rows for efficiency."""
    placeholders = []
    params = {"kyc": KYC_COMPLETED, "emp": EMP_ID, "assigned": ASSIGNED_TO}
    for i, name in enumerate(names):
        key = f"n{i}"
        placeholders.append(f"(%({key})s, %(kyc)s, %(emp)s, %(assigned)s)")
        params[key] = name

    query = (
        "INSERT INTO client (name, kyc_completed, emp_id, assigned_to) VALUES "
        + ", ".join(placeholders)
        + " RETURNING cli_id, name"
    )
    return query, params


def main():
    names = load_customer_names(EXCEL_PATH)
    print(f"Loaded {len(names)} customers from Excel.")

    if not names:
        print("No customers to insert.")
        return

    total_inserted = 0

    with psycopg.connect(CONN_STR) as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            for i in range(0, len(names), BATCH_SIZE):
                batch = names[i : i + BATCH_SIZE]
                query, params = build_batch_insert(batch)
                cur.execute(query, params)
                rows = cur.fetchall()
                total_inserted += len(rows)
                print(f"  Batch {i // BATCH_SIZE + 1}: inserted {len(rows)} customers")

        conn.commit()

    print(f"\nDone. Total inserted: {total_inserted}")


if __name__ == "__main__":
    main()
